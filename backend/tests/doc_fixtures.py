"""Builds real PDF and DOCX files for the extraction tests.

Genuine files rather than mocks: extraction is exactly the layer where a mock
would hide the parser behaviour under test.
"""

from __future__ import annotations

from pathlib import Path


def write_pdf(path: Path, pages: list[str]) -> Path:
    """Write a minimal but structurally valid PDF with one text line per page.

    A page given as an empty string has no text layer, which is how a scanned
    document looks to a text extractor.
    """
    objects: list[bytes] = []
    kids = " ".join(f"{4 + 2 * index} 0 R" for index in range(len(pages)))
    objects.append(b"<</Type/Catalog/Pages 2 0 R>>")
    objects.append(f"<</Type/Pages/Kids[{kids}]/Count {len(pages)}>>".encode())
    objects.append(b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>")

    for index, text in enumerate(pages):
        objects.append(
            (
                f"<</Type/Page/Parent 2 0 R/Resources<</Font<</F1 3 0 R>>>>"
                f"/MediaBox[0 0 612 792]/Contents {5 + 2 * index} 0 R>>"
            ).encode()
        )
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(
            b"<</Length " + str(len(stream)).encode() + b">>stream\n" + stream + b"\nendstream"
        )

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj".encode() + body + b"endobj\n"

    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer<</Size {len(objects) + 1}/Root 1 0 R>>\nstartxref\n{xref_at}\n%%EOF\n"
    ).encode()

    path.write_bytes(bytes(out))
    return path


def write_encrypted_pdf(
    path: Path,
    pages: list[str],
    password: str = "s3cret",  # noqa: S107 - the lock this fixture exists to create
) -> Path:
    """Write a PDF locked with a real user password.

    Encryption is applied by pypdf rather than faked, so the extractor meets the
    same refusal a genuinely protected file produces.
    """
    from pypdf import PdfReader, PdfWriter

    plain = path.with_name(f"{path.stem}-plain.pdf")
    write_pdf(plain, pages)
    try:
        writer = PdfWriter()
        for page in PdfReader(str(plain)).pages:
            writer.add_page(page)
        writer.encrypt(password)
        with path.open("wb") as handle:
            writer.write(handle)
    finally:
        plain.unlink(missing_ok=True)
    return path


def write_scanned_pdf(path: Path, page_images: list[bytes]) -> Path:
    """Write a PDF whose pages are images, with no text layer at all.

    This is what a scanner produces: extraction finds nothing and the pages have
    to be rendered and read as pixels.
    """
    import io

    from PIL import Image

    if not page_images:
        raise ValueError("a scanned PDF needs at least one page image")

    decoded = [Image.open(io.BytesIO(payload)).convert("RGB") for payload in page_images]
    first, *rest = decoded
    first.save(str(path), format="PDF", save_all=bool(rest), append_images=rest)
    return path


def write_docx(
    path: Path, paragraphs: list[str], table_rows: list[list[str]] | None = None
) -> Path:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for cell_index, value in enumerate(row):
                table.rows[row_index].cells[cell_index].text = value

    document.save(str(path))
    return path
