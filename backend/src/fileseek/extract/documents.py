"""Pulls readable text out of the document formats the library accepts.

A document with no text layer (a scan saved as PDF) is reported as such rather
than treated as empty, so the caller can fall back to character recognition.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Literal

TextSource = Literal["embedded", "ocr", "none"]

# Below this a "text layer" is really just stray artefacts from a scan.
MIN_MEANINGFUL_CHARS = 16


class UnreadableDocumentError(RuntimeError):
    """The file could not be parsed: encrypted, truncated, or not what it claims."""

    def __init__(self, path: Path, reason: str) -> None:
        super().__init__(f"could not read document {path.name}: {reason}")
        self.path = path
        self.reason = reason


@dataclass(frozen=True)
class ExtractedText:
    text: str
    page_count: int | None
    source: TextSource

    @property
    def has_text(self) -> bool:
        return len(self.text.strip()) >= MIN_MEANINGFUL_CHARS

    @property
    def needs_ocr(self) -> bool:
        """True for a document whose pages carry no usable text layer."""
        return self.source == "embedded" and not self.has_text


def _normalize(raw: str) -> str:
    """Collapse the whitespace noise that PDF and DOCX extraction leaves behind."""
    lines = [line.strip() for line in raw.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    kept: list[str] = []
    blank_run = 0
    for line in lines:
        if line:
            blank_run = 0
            kept.append(line)
            continue
        blank_run += 1
        # One blank line separates paragraphs; more is just layout residue.
        if blank_run == 1 and kept:
            kept.append("")
    while kept and not kept[-1]:
        kept.pop()
    return "\n".join(kept)


def _extract_plain_text(path: Path) -> ExtractedText:
    # utf-8-sig first: plain utf-8 decodes a BOM "successfully" but leaves the
    # marker in the text, where it would pollute the first chunk and the index.
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            raw = path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
        except OSError as error:
            raise UnreadableDocumentError(path, str(error)) from error
        return ExtractedText(text=_normalize(raw), page_count=None, source="embedded")
    raise UnreadableDocumentError(path, "could not decode the file as text")


def _unlock_pdf(reader: object, path: Path) -> None:
    """Open a PDF encrypted with an empty user password, which is common and harmless.

    A document with a real password stays closed and is reported as such.
    """
    try:
        opened = reader.decrypt("")  # type: ignore[attr-defined]
    except Exception as error:
        raise UnreadableDocumentError(path, f"the PDF is password protected ({error})") from error
    if opened == 0:
        raise UnreadableDocumentError(path, "the PDF is password protected")


def _read_pdf_pages(path: Path) -> list[str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        _unlock_pdf(reader, path)
    return [page.extract_text() or "" for page in reader.pages]


def _extract_pdf(path: Path) -> ExtractedText:
    from pypdf.errors import PdfReadError

    try:
        pages = _read_pdf_pages(path)
    except UnreadableDocumentError:
        raise
    except (PdfReadError, OSError, ValueError) as error:
        raise UnreadableDocumentError(path, str(error)) from error

    return ExtractedText(
        text=_normalize("\n".join(pages)), page_count=len(pages), source="embedded"
    )


def _extract_docx(path: Path) -> ExtractedText:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(str(path))
    except (PackageNotFoundError, OSError, ValueError, KeyError) as error:
        raise UnreadableDocumentError(path, str(error)) from error

    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text.strip() for cell in row.cells))

    # DOCX has no fixed pagination, so page_count stays unknown rather than guessed.
    return ExtractedText(text=_normalize("\n".join(blocks)), page_count=None, source="embedded")


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".txt": _extract_plain_text,
    ".md": _extract_plain_text,
}

# Enough resolution for character recognition without rendering print-size pages.
OCR_RENDER_SCALE = 2.0

# Recognition reads body text, so pages keep far more detail than the 512px the
# image encoder needs: a whole page shrunk to 512px is illegible.
OCR_MAX_EDGE = 2000

# A scan's first pages carry the title and abstract, which is what a search for
# the document is most likely to match; rendering all of a long scan is not worth it.
MAX_OCR_PAGES = 10


def rasterize_pages(
    path: Path, max_pages: int = MAX_OCR_PAGES, scale: float = OCR_RENDER_SCALE
) -> list[bytes]:
    """Render the first pages of a PDF to PNG bytes, for recognition.

    Only used when a PDF has no text layer: a scan has to be looked at as pixels.
    """
    import pypdfium2 as pdfium

    from fileseek.extract.images import encode_for_worker

    if max_pages < 1:
        raise ValueError(f"max_pages must be positive, got {max_pages}")

    try:
        document = pdfium.PdfDocument(str(path))
    except Exception as error:
        raise UnreadableDocumentError(path, f"could not open for rendering ({error})") from error

    try:
        rendered: list[bytes] = []
        for index in range(min(len(document), max_pages)):
            bitmap = document[index].render(scale=scale)
            rendered.append(encode_for_worker(bitmap.to_pil(), max_edge=OCR_MAX_EDGE))
    except Exception as error:
        raise UnreadableDocumentError(path, f"could not render pages ({error})") from error
    else:
        return rendered
    finally:
        document.close()


def extract_text(path: Path) -> ExtractedText:
    """Read a document's text, choosing the reader from its extension."""
    if not path.is_file():
        raise UnreadableDocumentError(path, "file does not exist")
    extractor = _EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        raise UnreadableDocumentError(path, f"unsupported document type {path.suffix!r}")
    return extractor(path)
